from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import sqlite3
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
import uuid
from datetime import datetime, timezone
from email.parser import BytesParser
from email.policy import default
from html.parser import HTMLParser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from sellersprite_connector import open_login as sellersprite_open_login
from sellersprite_connector import sync_results as sellersprite_sync_results
from sellersprite_connector import validate_sellersprite_url
from amazon_browser_connector import AmazonBrowserSession
from amazon_review_connector import collect_reviews as amazon_collect_reviews


ROOT = Path(__file__).resolve().parent
STATIC = (ROOT / "static").resolve()
DEFAULT_DATA = Path(r"D:\AM-master-data") if os.name == "nt" else ROOT / "data"
DATA_DIR = Path(os.environ.get("AM_MASTER_DATA_DIR", str(DEFAULT_DATA))).expanduser()
DB_PATH = DATA_DIR / "am_master.sqlite3"
RAW_DIR = DATA_DIR / "raw"
AMAZON_CACHE_DIR = DATA_DIR / "cache" / "amazon"
AMAZON_PROFILE_DIR = DATA_DIR / "browser-profile" / "amazon"
SELLERSPRITE_PROFILE_DIR = DATA_DIR / "browser-profile" / "sellersprite"
PORT = int(os.environ.get("AM_MASTER_PORT", "8777"))
MAX_PRODUCTS_PER_MARKET = 100
MAX_SEARCH_PAGES = 10
AMAZON_FETCH_RETRIES = max(1, int(os.environ.get("AM_MASTER_AMAZON_RETRIES", "4")))
AMAZON_BROWSER_RECOVERY = os.environ.get("AM_MASTER_AMAZON_BROWSER_RECOVERY", "1") != "0"
_DB_INIT_LOCK = threading.Lock()
_INITIALIZED_DATABASES = set()
_CONNECTOR_LOCK = threading.Lock()
_CONNECTOR_RUNS = {}
COUNTRIES = {
    "US": ("美国", "www.amazon.com", "USD", "en-US,en;q=0.9"),
    "CA": ("加拿大", "www.amazon.ca", "CAD", "en-CA,en;q=0.9"),
    "UK": ("英国", "www.amazon.co.uk", "GBP", "en-GB,en;q=0.9"),
    "DE": ("德国", "www.amazon.de", "EUR", "de-DE,de;q=0.9"),
    "FR": ("法国", "www.amazon.fr", "EUR", "fr-FR,fr;q=0.9"),
    "IT": ("意大利", "www.amazon.it", "EUR", "it-IT,it;q=0.9"),
    "ES": ("西班牙", "www.amazon.es", "EUR", "es-ES,es;q=0.9"),
}
FEATURES = {
    "双仓/多仓": ["双粮仓", "双仓", "dual hopper", "two hopper", "dual chamber", "2 hopper"],
    "摄像头": ["摄像头", "摄像", "camera", "video"],
    "Wi-Fi/APP": ["wifi", "wi-fi", "app", "手机控制", "远程控制"],
    "定时定量": ["定时", "定量", "计划喂养", "schedule", "portion", "timed", "programmable"],
    "克重校准": ["克重", "克数", "重量校准", "weigh", "gram", "scale"],
    "备用电源": ["备用电池", "断电", "dual power", "battery backup", "battery operated"],
    "防堵/堵粮提醒": ["防堵", "堵粮", "anti-jam", "anti clog", "blockage sensor"],
    "密封保鲜": ["密封", "保鲜", "干燥剂", "airtight", "freshness", "desiccant", "sealed"],
    "不锈钢餐盘": ["不锈钢", "304", "316", "stainless steel"],
    "双向语音": ["对讲", "双向语音", "two-way audio", "2-way audio"],
    "夜视": ["夜视", "红外", "night vision", "infrared", "940nm"],
    "可拆洗": ["可拆洗", "可拆卸", "易清洁", "detachable", "dishwasher", "easy clean"],
    "指示灯控制": ["指示灯", "灯光关闭", "indicator light", "light control"],
    "多宠识别": ["宠物识别", "rfid", "microchip", "pet recognition", "multi-pet"],
}

REVIEW_THEMES = {
    "出粮可靠性": ["jam", "stuck", "clog", "didn't dispense", "not dispense", "failed to feed", "堵", "卡粮", "不出粮", "bloqué", "ne distribue", "non eroga", "atascado", "no dispensa", "klemmt", "gibt kein futter"],
    "份量准确": ["portion", "too much", "too little", "inaccurate", "grams", "份量", "克重", "太多", "太少", "portion", "quantité", "porzione", "cantidad", "menge"],
    "APP与连接": ["app", "wifi", "wi-fi", "connect", "offline", "timeout", "connexion", "verbindung", "connessione", "conexión", "连接", "离线", "超时"],
    "摄像头与夜视": ["camera", "video", "night vision", "blurry", "pixel", "caméra", "vision nocturne", "kamera", "nachtsicht", "telecamera", "visione notturna", "cámara", "visión nocturna", "摄像", "夜视", "模糊"],
    "语音与噪声": ["audio", "speaker", "microphone", "noise", "echo", "squeal", "son", "lautsprecher", "geräusch", "audio", "rumore", "sonido", "ruido", "对讲", "啸叫", "噪声"],
    "密封与保鲜": ["seal", "fresh", "moisture", "airtight", "desiccant", "dichtung", "frisch", "étanche", "fraîcheur", "sigillo", "fresco", "sellado", "fresco", "密封", "保鲜", "受潮"],
    "清洁与材质": ["clean", "wash", "stainless", "plastic", "dishwasher", "reinigen", "edelstahl", "nettoyer", "inox", "pulire", "acciaio", "limpiar", "acero", "清洗", "不锈钢", "塑料"],
    "电源与续航": ["battery", "power", "outage", "backup", "strom", "batterie", "courant", "batteria", "energía", "电池", "断电", "备用电源"],
    "误触与提醒": ["false alert", "notification", "sensor", "accidental", "touch", "fehlalarm", "alerte", "notifica", "alerta", "误触", "误报", "提醒", "传感"],
    "安全与防夹": ["unsafe", "danger", "pinch", "trap", "stuck inside", "child lock", "anti-pinch", "safety sensor", "gefährlich", "eingeklemmt", "sécurité", "coincé", "sicurezza", "incastrato", "seguridad", "atrapado", "防夹", "安全", "夹猫"],
    "清洁循环与卡机": ["cleaning cycle", "cycle failed", "won't rotate", "doesn't rotate", "jammed", "stops cycling", "reinigungszyklus", "dreht nicht", "cycle de nettoyage", "ne tourne pas", "ciclo di pulizia", "non gira", "ciclo de limpieza", "no gira", "清洁循环", "不转", "卡机"],
    "漏砂与带砂": ["litter leak", "leaking litter", "tracks litter", "litter everywhere", "streu", "litière partout", "sabbia fuori", "arena fuera", "漏砂", "带砂", "猫砂到处"],
    "除臭与集便仓": ["odor", "odour", "smell", "waste drawer", "waste bin", "full bin", "geruch", "abfallbehälter", "odeur", "bac à déchets", "odore", "cassetto rifiuti", "olor", "cajón de residuos", "臭味", "异味", "集便仓", "垃圾仓"],
    "猫砂兼容": ["litter type", "crystal litter", "clumping litter", "pellet", "tofu litter", "streuart", "type de litière", "tipo di lettiera", "tipo de arena", "猫砂兼容", "豆腐砂", "膨润土", "水晶砂"],
    "多猫与体重识别": ["multiple cats", "multi-cat", "weight tracking", "weight monitor", "can't tell cats", "mehrere katzen", "poids", "più gatti", "peso", "varios gatos", "多猫", "体重识别", "分不清猫"],
    "耗材与维护成本": ["bag", "liner", "filter", "deodorizer", "subscription", "replacement", "beutel", "sac", "filtro", "bolsa", "耗材", "垃圾袋", "滤芯", "除臭盒"],
}

LITTER_REVIEW_THEME_NAMES = {
    "APP与连接", "清洁与材质", "电源与续航", "误触与提醒", "安全与防夹",
    "清洁循环与卡机", "漏砂与带砂", "除臭与集便仓", "猫砂兼容",
    "多猫与体重识别", "耗材与维护成本", "摄像头与夜视", "语音与噪声",
}

LITTER_ACCESSORY_TERMS = (
    "bag", "liner", "mat", "carpet", "filter", "deodorizer", "replacement",
    "accessory", "scoop", "refill", "waste bags", "deodorization box",
    "litter box enclosure", "litter box furniture", "litter box cabinet",
    "litter box monitor", "sac", "tapis", "filtre",
    "rechange", "pelle", "meuble litière", "beutel", "matte", "ersatz",
    "schaufel", "katzentoilettenmöbel", "filtro", "ricambio", "sacchetto",
    "paletta", "mobile lettiera", "bolsa", "alfombra", "repuesto", "pala",
    "mueble arenero", "垃圾袋", "垫子", "滤芯", "替换件",
    "除臭盒", "猫砂铲", "猫砂盆家具",
)
LITTER_STRONG_ACCESSORY_TERMS = (
    "replacement motor", "replacement liner", "replacement part", "replacement waste",
    "waste bags for", "litter bags for", "liners for", "liners fit", "refill ring",
    "refill for", "compatible with", "litter box enclosure", "litter box furniture",
    "litter box cabinet", "litter box monitor", "deodorization box", "motor for",
    "door curtain", "base mat", "waterproof pad", "ramp with filter", "filter for",
    "special filter", "pack refill", "beutel für", "ersatzmotor", "ersatzteil",
    "sacs pour", "recharge pour", "tapis de base", "filtro per", "ricambio per",
    "bolsas para", "repuesto para", "motor de repuesto", "配件", "替换电机",
    "替换内衬", "专用垃圾袋", "家具柜", "监测器", "tray refills",
    "litter tray bags", "automatic litter box bags", "self-cleaning litter bags",
    "cat litter bag with", "supply & accessory bundle", "accessory bundle",
    "air purifying refill", "cat litter box liners", "waste drawer liners",
)
LITTER_AUTOMATION_TERMS = (
    "self-cleaning", "self cleaning", "automatic cleaning", "automatic litter",
    "automatic cat toilet", "auto-clean", "autoclean", "litter robot", "litter-robot",
    "robotic litter", "smart litter", "electric litter", "wifi litter",
    "selbstreinig", "automatische katzentoilette", "smartes katzenklo",
    "litière autonettoyante", "litiere autonettoyante", "litière automatique",
    "bac à litière automatique", "lettiera autopulente", "lettiera automatica",
    "arenero autolimpiante", "arenero autolimpiable", "arenero automático",
    "caja de arena automática", "caja de arena autolimpiante",
    "caja de arena autolimpiable", "bandeja para gatos de limpieza automática",
    "bandeja de arena para gatos autolimpiable", "arena para gatos autolimpiante",
    "inodoro autolimpiable para gatos", "arenero rotatorio autolimpiable",
    "自动猫砂盆", "智能猫砂盆", "自清洁猫砂盆",
)
LITTER_MANUAL_TERMS = (
    "manual self-cleaning", "roll 'n clean", "roll n clean", "roll'n clean",
    "roll and clean", "rolling litter box", "sift and clean", "omega paw",
    "catit smartsift", "bandeja sanitaria automática", "手动滚筒",
)


def classify_litter_product_title(title):
    text = re.sub(r"\s+", " ", str(title or "").lower()).strip()
    def matches(term):
        if term in {"selbstreinig"}:
            return term in text
        if re.fullmatch(r"[a-z0-9 -]+", term):
            return re.search(rf"(?<![a-z0-9]){re.escape(term)}(?![a-z0-9])", text) is not None
        return term in text

    accessory_hits = [term for term in LITTER_ACCESSORY_TERMS if matches(term)]
    strong_accessory_hits = [term for term in LITTER_STRONG_ACCESSORY_TERMS if matches(term)]
    automation_hits = [term for term in LITTER_AUTOMATION_TERMS if matches(term)]
    manual_hits = [term for term in LITTER_MANUAL_TERMS if matches(term)]
    if manual_hits:
        return {"status": "excluded_manual", "reason": manual_hits[0]}
    counted_accessory = bool(accessory_hits and re.match(r"^\s*\d+\s*(?:count|pcs|pack|stück|pièces|pezzi|unidades)?\b", text))
    if strong_accessory_hits or counted_accessory:
        return {"status": "excluded_accessory", "reason": strong_accessory_hits[0] if strong_accessory_hits else accessory_hits[0]}
    if automation_hits:
        return {"status": "qualified", "reason": automation_hits[0]}
    if accessory_hits:
        return {"status": "excluded_accessory", "reason": accessory_hits[0]}
    return {"status": "unconfirmed", "reason": "title lacks explicit automation evidence"}


def connect():
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(DB_PATH, timeout=30)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA journal_mode=WAL")
    database_id = str(DB_PATH.resolve())
    if database_id not in _INITIALIZED_DATABASES:
        with _DB_INIT_LOCK:
            if database_id not in _INITIALIZED_DATABASES:
                con.executescript("""
                  CREATE TABLE IF NOT EXISTS products(
                    id INTEGER PRIMARY KEY, country TEXT NOT NULL, category_path TEXT NOT NULL DEFAULT '',
                    source TEXT NOT NULL, source_date TEXT NOT NULL, asin TEXT NOT NULL DEFAULT '',
                    title TEXT NOT NULL, brand TEXT NOT NULL DEFAULT '', price REAL, currency TEXT,
                    units REAL, revenue REAL, growth REAL, rating REAL, reviews REAL, bsr_rank REAL,
                    image_url TEXT, product_url TEXT, product_key TEXT NOT NULL DEFAULT '',
                    raw_json TEXT NOT NULL DEFAULT '{}',
                    UNIQUE(country, source, source_date, asin, title)
                  );
                  CREATE INDEX IF NOT EXISTS ix_products_scope ON products(country,category_path,source_date);
                  CREATE TABLE IF NOT EXISTS searches(
                    id TEXT PRIMARY KEY, keyword TEXT NOT NULL, category_path TEXT NOT NULL,
                    countries TEXT NOT NULL, status TEXT NOT NULL, started_at TEXT NOT NULL,
                    finished_at TEXT, message TEXT NOT NULL DEFAULT ''
                  );
                  CREATE TABLE IF NOT EXISTS analyses(
                    id TEXT PRIMARY KEY, product_name TEXT NOT NULL, category_path TEXT NOT NULL,
                    product_text TEXT NOT NULL, created_at TEXT NOT NULL, result_json TEXT NOT NULL
                  );
                  CREATE TABLE IF NOT EXISTS connector_runs(
                    id TEXT PRIMARY KEY, kind TEXT NOT NULL, status TEXT NOT NULL,
                    started_at TEXT NOT NULL, finished_at TEXT, message TEXT NOT NULL DEFAULT '',
                    result_json TEXT NOT NULL DEFAULT '{}'
                  );
                  CREATE TABLE IF NOT EXISTS amazon_search_cache(
                    country TEXT NOT NULL, keyword TEXT NOT NULL, category_path TEXT NOT NULL,
                    products_json TEXT NOT NULL, updated_at TEXT NOT NULL,
                    PRIMARY KEY(country,keyword,category_path)
                  );
                  CREATE TABLE IF NOT EXISTS reviews(
                    id INTEGER PRIMARY KEY, country TEXT NOT NULL, asin TEXT NOT NULL,
                    product_title TEXT NOT NULL DEFAULT '', review_id TEXT NOT NULL,
                    rating REAL, review_title TEXT NOT NULL DEFAULT '', review_body TEXT NOT NULL,
                    review_date_text TEXT NOT NULL DEFAULT '', verified INTEGER NOT NULL DEFAULT 0,
                    helpful_count INTEGER NOT NULL DEFAULT 0, reviewer TEXT NOT NULL DEFAULT '',
                    review_url TEXT NOT NULL DEFAULT '', source TEXT NOT NULL,
                    collected_at TEXT NOT NULL, themes_json TEXT NOT NULL DEFAULT '[]',
                    UNIQUE(country,asin,review_id)
                  );
                  CREATE INDEX IF NOT EXISTS ix_reviews_scope ON reviews(country,asin,collected_at);
                  CREATE TABLE IF NOT EXISTS analysis_runs(
                    id TEXT PRIMARY KEY, kind TEXT NOT NULL, product_name TEXT NOT NULL DEFAULT '',
                    category_path TEXT NOT NULL DEFAULT '', countries TEXT NOT NULL DEFAULT '[]',
                    source_period TEXT NOT NULL DEFAULT '', sample_count INTEGER NOT NULL DEFAULT 0,
                    started_at TEXT NOT NULL, finished_at TEXT, status TEXT NOT NULL,
                    conclusions_json TEXT NOT NULL DEFAULT '{}', message TEXT NOT NULL DEFAULT ''
                  );
                """)
                columns = {row[1] for row in con.execute("PRAGMA table_info(products)")}
                if "product_key" not in columns:
                    con.execute("ALTER TABLE products ADD COLUMN product_key TEXT NOT NULL DEFAULT ''")
                if "bsr_rank" not in columns:
                    con.execute("ALTER TABLE products ADD COLUMN bsr_rank REAL")
                migrate_bsr_ranks(con)
                migrate_product_keys(con)
                con.execute("CREATE UNIQUE INDEX IF NOT EXISTS ux_products_country_key ON products(country,product_key) WHERE product_key<>''")
                con.commit()
                _INITIALIZED_DATABASES.add(database_id)
    return con


def clean_number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    text = re.sub(r"[^0-9,.\-]", "", text)
    if not text:
        return None
    if "," in text and "." in text:
        text = text.replace(",", "") if text.rfind(".") > text.rfind(",") else text.replace(".", "").replace(",", ".")
    elif "," in text:
        tail = text.split(",")[-1]
        text = text.replace(",", "." if len(tail) in (1, 2) else "")
    try:
        return float(text)
    except ValueError:
        return None


def clean_compact_number(value):
    if value is None:
        return None
    text = str(value).strip().upper().replace(" ", "")
    multiplier = 1
    if text.endswith("MILA"):
        multiplier = 1_000
        text = text[:-4]
    elif text.endswith("MIL"):
        multiplier = 1_000
        text = text[:-3]
    elif text.endswith("K"):
        multiplier = 1_000
        text = text[:-1]
    elif text.endswith("M"):
        multiplier = 1_000_000
        text = text[:-1]
    number = clean_number(text)
    return number * multiplier if number is not None else None


def clean_first_number(value, compact=False):
    match = re.search(r"-?\d[\d.,]*(?:\s*(?:MILA|MIL|K|M))?" if compact else r"-?\d[\d.,]*", str(value or ""), re.I)
    if not match:
        return None
    return clean_compact_number(match.group(0)) if compact else clean_number(match.group(0))


def currency_from_text(value, base_url=""):
    text = str(value or "").upper().replace(" ", "")
    if "CNY" in text or "CN¥" in text or "RMB" in text:
        return "CNY"
    if "CAD" in text or "CA$" in text or "C$" in text:
        return "CAD"
    if "USD" in text or "US$" in text:
        return "USD"
    if "GBP" in text or "£" in text:
        return "GBP"
    if "EUR" in text or "€" in text:
        return "EUR"
    host = urllib.parse.urlparse(base_url).hostname or ""
    if "$" in text:
        return "CAD" if host.endswith("amazon.ca") else "USD"
    if host.endswith("amazon.co.uk"):
        return "GBP"
    if any(host.endswith(domain) for domain in ("amazon.de", "amazon.fr", "amazon.it", "amazon.es")):
        return "EUR"
    return "USD" if host.endswith("amazon.com") else ""


def normalized_product_url(value):
    try:
        parsed = urllib.parse.urlparse(str(value or "").strip())
    except ValueError:
        return ""
    if parsed.scheme not in ("http", "https") or not parsed.netloc:
        return ""
    path = re.sub(r"/(ref|gp/aw/d)/.*$", "", parsed.path.rstrip("/"), flags=re.I)
    return f"{parsed.netloc.lower()}{path.lower()}"


def product_identity(row):
    asin = re.sub(r"[^A-Z0-9]", "", str(row.get("asin") or "").upper())
    if asin:
        return f"asin:{asin}"
    match = re.search(r"/(?:dp|gp/product)/([A-Z0-9]{10})(?:[/?]|$)", str(row.get("product_url") or ""), re.I)
    if match:
        return f"asin:{match.group(1).upper()}"
    url = normalized_product_url(row.get("product_url"))
    if url:
        return f"url:{url}"
    title = re.sub(r"[^a-z0-9\u4e00-\u9fff]", "", str(row.get("title") or "").lower())
    return f"title:{title}" if title else ""


def deduplicate_products(rows, limit=None):
    unique = []
    seen = set()
    for row in rows:
        key = product_identity(row)
        if not key or key in seen:
            continue
        seen.add(key)
        unique.append(row)
        if limit and len(unique) >= limit:
            break
    return unique


def describe_request_error(exc):
    if isinstance(exc, urllib.error.HTTPError):
        explanations = {
            403: "Amazon 拒绝访问（HTTP 403）",
            429: "请求频率受限（HTTP 429）",
            503: "Amazon 临时拒绝或服务不可用（HTTP 503）",
        }
        return explanations.get(exc.code, f"Amazon 请求失败（HTTP {exc.code}）")
    if isinstance(exc, TimeoutError):
        return "请求超时"
    if isinstance(exc, urllib.error.URLError):
        return f"网络连接失败：{exc.reason}"
    return str(exc) or exc.__class__.__name__


def detect_page_issue(raw):
    if not raw:
        return "Amazon 返回空页面"
    sample = raw[:20000].lower()
    if b"bm-verify" in sample or b"interstitialchallenge" in sample:
        return "Amazon 返回风控验证页（bm-verify），未绕过验证"
    if b"captcha" in sample or b"robot check" in sample:
        return "Amazon 返回验证码页面，未绕过验证"
    return "页面可访问，但没有识别到商品；可能是页面结构变化或非商品搜索页"


def migrate_bsr_ranks(con):
    rows = con.execute("SELECT id,source,raw_json FROM products WHERE bsr_rank IS NULL AND source LIKE 'SellerSprite%'").fetchall()
    for row in rows:
        try:
            payload = json.loads(row["raw_json"] or "{}")
        except (ValueError, TypeError):
            continue
        rank = clean_number(payload.get("bsr_rank", payload.get("rank")))
        if rank is not None:
            con.execute("UPDATE products SET bsr_rank=? WHERE id=?", (rank, row["id"]))


def amazon_page_cache_path(country, keyword, page):
    digest = hashlib.sha256(f"{country}\n{keyword.strip().lower()}\n{page}".encode("utf-8")).hexdigest()[:24]
    return AMAZON_CACHE_DIR / country / f"{digest}.json"


def save_amazon_page_cache(country, keyword, page, url, raw):
    path = amazon_page_cache_path(country, keyword, page)
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "country": country, "keyword": keyword, "page": page, "url": url,
        "fetched_at": datetime.now(timezone.utc).isoformat(),
        "html": raw.decode("utf-8", "replace"),
    }
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")


def load_amazon_page_cache(country, keyword, page):
    path = amazon_page_cache_path(country, keyword, page)
    if not path.exists():
        return None
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
        raw = str(payload.get("html") or "").encode("utf-8")
        return (raw, payload.get("fetched_at", "")) if raw else None
    except (OSError, ValueError, TypeError):
        return None


def load_search_snapshot(country, keyword, category_path):
    con = connect()
    row = con.execute(
        "SELECT products_json,updated_at FROM amazon_search_cache WHERE country=? AND keyword=? AND category_path=?",
        (country, keyword.strip().lower(), category_path.strip()),
    ).fetchone()
    con.close()
    if not row:
        return [], ""
    try:
        return json.loads(row["products_json"]), row["updated_at"]
    except (ValueError, TypeError):
        return [], ""


def save_search_snapshot(country, keyword, category_path, rows):
    rows = deduplicate_products(rows, MAX_PRODUCTS_PER_MARKET)
    if not rows:
        return
    con = connect()
    existing = con.execute(
        "SELECT products_json FROM amazon_search_cache WHERE country=? AND keyword=? AND category_path=?",
        (country, keyword.strip().lower(), category_path.strip()),
    ).fetchone()
    existing_count = 0
    if existing:
        try:
            existing_count = len(json.loads(existing["products_json"]))
        except (ValueError, TypeError):
            pass
    if len(rows) >= existing_count:
        con.execute(
            "INSERT INTO amazon_search_cache(country,keyword,category_path,products_json,updated_at) VALUES(?,?,?,?,?) "
            "ON CONFLICT(country,keyword,category_path) DO UPDATE SET products_json=excluded.products_json,updated_at=excluded.updated_at",
            (country, keyword.strip().lower(), category_path.strip(), json.dumps(rows, ensure_ascii=False), datetime.now(timezone.utc).isoformat()),
        )
        con.commit()
    con.close()


def load_bsr_category_products(country, category_path, limit=MAX_PRODUCTS_PER_MARKET):
    if not category_path.strip():
        return []
    con = connect()
    rows = [dict(row) for row in con.execute(
        "SELECT * FROM products WHERE country=? AND category_path=? "
        "AND bsr_rank IS NOT NULL "
        "ORDER BY bsr_rank ASC, COALESCE(units,0) DESC, COALESCE(revenue,0) DESC, source_date DESC LIMIT ?",
        (country, category_path.strip(), int(limit)),
    )]
    con.close()
    return deduplicate_products(rows, limit)


def merge_bsr_then_search(bsr_rows, search_rows, limit=MAX_PRODUCTS_PER_MARKET):
    ranked = sorted((dict(row) for row in bsr_rows), key=lambda row: (row.get("bsr_rank") is None, row.get("bsr_rank") or float("inf")))
    return deduplicate_products([*ranked, *search_rows], limit)


def migrate_product_keys(con):
    rows = [dict(row) for row in con.execute("SELECT * FROM products ORDER BY source_date DESC, id DESC")]
    keepers = {}
    for row in rows:
        key = product_identity(row)
        if not key:
            continue
        keeper = keepers.get((row["country"], key))
        if keeper is None:
            keepers[(row["country"], key)] = row
            con.execute("UPDATE products SET product_key=? WHERE id=?", (key, row["id"]))
            continue
        updates = {}
        keeper_has_metrics = any(keeper.get(field) is not None for field in ("units", "revenue", "growth"))
        row_has_metrics = any(row.get(field) is not None for field in ("units", "revenue", "growth"))
        if row_has_metrics and not keeper_has_metrics:
            updates["source"] = row["source"]
            keeper["source"] = row["source"]
        for field in ("units", "revenue", "growth", "price", "rating", "reviews", "bsr_rank", "image_url", "product_url"):
            if keeper.get(field) in (None, "") and row.get(field) not in (None, ""):
                updates[field] = row[field]
                keeper[field] = row[field]
        if updates:
            assignment = ",".join(f"{field}=?" for field in updates)
            con.execute(f"UPDATE products SET {assignment} WHERE id=?", (*updates.values(), keeper["id"]))
        con.execute("DELETE FROM products WHERE id=?", (row["id"],))


def normalize_header(value):
    text = str(value or "").strip().lower()
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]", "", text)


HEADER_ALIASES = {
    "country": ["country", "国家", "站点", "marketplace", "market"],
    "category_path": ["category", "categorypath", "品类", "三级类目", "类目路径"],
    "asin": ["asin", "productid", "商品编码", "产品id"],
    "title": ["title", "商品标题", "产品标题", "商品名称", "产品名称"],
    "brand": ["brand", "品牌"], "price": ["price", "价格", "售价"],
    "currency": ["currency", "币种"], "units": ["units", "销量", "月销量", "月销售量"],
    "revenue": ["revenue", "销售额", "月销售额", "月销额"],
    "growth": ["growth", "增长率", "销量增长", "销售额增长", "环比增长"],
    "rating": ["rating", "评分", "星级"], "reviews": ["reviews", "评论数", "评分数"],
    "bsr_rank": ["bsr", "bsrrank", "bsr排名", "大类排名", "类目排名", "排名"],
    "image_url": ["image", "imageurl", "图片", "图片链接"],
    "product_url": ["url", "producturl", "商品链接", "产品链接"],
}


def canonical_country(value):
    text = str(value or "").strip().upper()
    reverse = {label.upper(): code for code, (label, *_rest) in COUNTRIES.items()}
    for code in COUNTRIES:
        if text == code or text == COUNTRIES[code][0].upper() or text == COUNTRIES[code][1].upper():
            return code
    return reverse.get(text, "")


class AmazonSearchParser(HTMLParser):
    def __init__(self, base_url):
        super().__init__(convert_charrefs=True)
        self.base_url = base_url
        self.items = []
        self.item = None
        self.root_depth = 0
        self.depth = 0
        self.capture = None
        self.capture_depth = 0
        self.title_depth = None

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if self.item is None and tag == "div" and attrs.get("data-asin") and attrs.get("data-component-type") == "s-search-result":
            self.item = {"asin": attrs["data-asin"], "title": "", "brand": "", "price": None, "currency": "", "rating": None, "reviews": None, "image_url": "", "product_url": ""}
            self.root_depth = self.depth
        if self.item is not None:
            cls = attrs.get("class", "")
            if tag == "h2":
                self.title_depth = self.depth
                if attrs.get("aria-label"):
                    self.item["title"] = attrs["aria-label"].strip()
            if tag == "img" and not self.item["image_url"]:
                self.item["image_url"] = attrs.get("src", "") or attrs.get("data-src", "")
            if tag == "a" and ("s-link-style" in cls or re.search(r"/dp/[A-Z0-9]{10}", attrs.get("href", ""))) and not self.item["product_url"]:
                self.item["product_url"] = urllib.parse.urljoin(self.base_url, attrs.get("href", ""))
            if tag == "span":
                if "a-offscreen" in cls and self.item["price"] is None: self.capture = "price_offscreen"
                elif "a-price-whole" in cls and self.item["price"] is None: self.capture = "price"
                elif "a-icon-alt" in cls and self.item["rating"] is None: self.capture = "rating"
                elif "s-underline-text" in cls and self.item["reviews"] is None: self.capture = "reviews"
                elif (self.title_depth is not None or "a-size-base-plus" in cls or "a-size-medium" in cls) and not self.item["title"]: self.capture = "title"
                elif "a-size-base" in cls and not self.item["brand"]: self.capture = "brand"
                if self.capture: self.capture_depth = self.depth
        if tag not in {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "track", "wbr"}:
            self.depth += 1

    def handle_endtag(self, tag):
        self.depth = max(0, self.depth - 1)
        if self.item is not None and self.depth <= self.root_depth:
            if self.item.get("title"):
                self.items.append(self.item)
            self.item = None
            self.capture = None
            self.title_depth = None
        elif self.capture and self.depth <= self.capture_depth:
            self.capture = None
        if self.title_depth is not None and self.depth <= self.title_depth:
            self.title_depth = None

    def handle_data(self, data):
        if self.item is None or not self.capture:
            return
        value = data.strip()
        if not value:
            return
        key = self.capture
        if key in ("price", "price_offscreen", "rating", "reviews"):
            self.item[key + "_raw"] = self.item.get(key + "_raw", "") + value
            raw = self.item[key + "_raw"]
            if key == "price_offscreen":
                self.item["price"] = clean_first_number(raw)
                self.item["currency"] = currency_from_text(raw, self.base_url)
            else:
                self.item[key] = clean_first_number(raw, compact=(key == "reviews"))
        else:
            self.item[key] = (self.item.get(key, "") + " " + value).strip()


def amazon_search(country, keyword, category_path, progress=None):
    label, host, currency, language = COUNTRIES[country]
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    folder = RAW_DIR / stamp
    folder.mkdir(parents=True, exist_ok=True)
    collected = []
    seen = set()
    final_url = ""
    empty_pages = 0
    browser = None
    terminal_issue = ""
    try:
        for page in range(1, MAX_SEARCH_PAGES + 1):
            query = urllib.parse.urlencode({"k": keyword, "page": page})
            request = urllib.request.Request(f"https://{host}/s?{query}", headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/130 Safari/537.36", "Accept-Language": language})
            raw = None
            issue = ""
            from_cache = False
            cache_date = ""
            for attempt in range(1, AMAZON_FETCH_RETRIES + 1):
                try:
                    with urllib.request.urlopen(request, timeout=25) as response:
                        candidate = response.read()
                        final_url = response.geturl()
                    candidate_parser = AmazonSearchParser(f"https://{host}")
                    candidate_parser.feed(candidate.decode("utf-8", "replace"))
                    if candidate_parser.items:
                        raw = candidate
                        issue = ""
                        save_amazon_page_cache(country, keyword, page, final_url, raw)
                        break
                    issue = detect_page_issue(candidate)
                except Exception as exc:
                    issue = describe_request_error(exc)
                if progress:
                    progress(page, len(collected), 0, f"{issue}；自动重试 {attempt}/{AMAZON_FETCH_RETRIES}")
                if attempt < AMAZON_FETCH_RETRIES:
                    time.sleep(min(8.0, 1.25 * (2 ** (attempt - 1))))

            if raw is None and AMAZON_BROWSER_RECOVERY:
                try:
                    if browser is None:
                        if progress:
                            progress(page, len(collected), 0, "常规请求未恢复，正在打开可见 Edge 浏览器")
                        browser = AmazonBrowserSession(AMAZON_PROFILE_DIR)
                    browser_progress = None
                    if progress:
                        browser_progress = lambda current_page, _count, added, warning: progress(current_page, len(collected), added, warning)
                    verification_expected = "验证" in issue or "captcha" in issue.lower() or "bm-verify" in issue.lower()
                    raw = browser.fetch(request.full_url, browser_progress, page, verification_expected=verification_expected)
                    if raw:
                        final_url = request.full_url
                        save_amazon_page_cache(country, keyword, page, final_url, raw)
                        issue = ""
                except Exception as exc:
                    issue = f"浏览器恢复失败：{exc}"

            if raw is None:
                cached = load_amazon_page_cache(country, keyword, page)
                if cached:
                    raw, cache_date = cached
                    from_cache = True
                    issue = f"实时页面未恢复，使用本机页面缓存（{cache_date[:10] or '历史'}）"

            if raw is None:
                terminal_issue = issue or "Amazon 页面未返回"
                if progress:
                    progress(page, len(collected), 0, terminal_issue)
                break

            (folder / f"amazon_{country}_page_{page}.html").write_bytes(raw)
            parser = AmazonSearchParser(f"https://{host}")
            parser.feed(raw.decode("utf-8", "replace"))
            if not parser.items:
                terminal_issue = detect_page_issue(raw)
                if progress:
                    progress(page, len(collected), 0, terminal_issue)
                break
            added = 0
            for item in parser.items:
                row = {"country": country, "category_path": category_path, "source": "Amazon cached public search" if from_cache else "Amazon public search", "source_date": cache_date or datetime.now(timezone.utc).isoformat(), "asin": item["asin"], "title": item["title"], "brand": item.get("brand", ""), "price": item.get("price"), "currency": item.get("currency") or currency, "units": None, "revenue": None, "growth": None, "rating": item.get("rating"), "reviews": item.get("reviews"), "image_url": item.get("image_url", ""), "product_url": item.get("product_url", "") or f"https://{host}/dp/{item['asin']}"}
                key = product_identity(row)
                if not key or key in seen:
                    continue
                seen.add(key)
                row["rank"] = len(collected) + 1
                collected.append(row)
                added += 1
                if len(collected) >= MAX_PRODUCTS_PER_MARKET:
                    break
            if progress:
                progress(page, len(collected), added, issue if from_cache else None)
            if len(collected) >= MAX_PRODUCTS_PER_MARKET:
                break
            empty_pages = empty_pages + 1 if added == 0 else 0
            if empty_pages >= 2:
                terminal_issue = "连续两页没有新增去重商品"
                break
            time.sleep(0.8)
    finally:
        if browser is not None:
            browser.close()

    if len(collected) < MAX_PRODUCTS_PER_MARKET:
        snapshot, snapshot_date = load_search_snapshot(country, keyword, category_path)
        added_from_snapshot = 0
        for row in snapshot:
            key = product_identity(row)
            if not key or key in seen:
                continue
            cached_row = dict(row)
            cached_row["source"] = "Amazon cached search snapshot"
            cached_row["source_date"] = snapshot_date or cached_row.get("source_date", "")
            cached_row["rank"] = len(collected) + 1
            collected.append(cached_row)
            seen.add(key)
            added_from_snapshot += 1
            if len(collected) >= MAX_PRODUCTS_PER_MARKET:
                break
        if added_from_snapshot and progress:
            progress(0, len(collected), added_from_snapshot, f"实时采集有缺口，已从本机查询快照补回 {added_from_snapshot} 条（{snapshot_date[:10] or '历史'}）")
    items = collected[:MAX_PRODUCTS_PER_MARKET]
    if not items:
        raise RuntimeError(f"没有取得商品。{terminal_issue}。最后访问地址：{final_url}")
    save_search_snapshot(country, keyword, category_path, items)
    return items


def save_products(items):
    items = deduplicate_products(items)
    if not items:
        return 0
    con = connect()
    for row in items:
        row = {**row, "bsr_rank": clean_number(row.get("bsr_rank")), "product_key": product_identity(row), "raw_json": json.dumps(row, ensure_ascii=False)}
        con.execute("""INSERT INTO products(country,category_path,source,source_date,asin,title,brand,price,currency,units,revenue,growth,rating,reviews,bsr_rank,image_url,product_url,product_key,raw_json)
          VALUES(:country,:category_path,:source,:source_date,:asin,:title,:brand,:price,:currency,:units,:revenue,:growth,:rating,:reviews,:bsr_rank,:image_url,:product_url,:product_key,:raw_json)
          ON CONFLICT(country,product_key) WHERE product_key<>'' DO UPDATE SET
            category_path=excluded.category_path,
            source=CASE
              WHEN excluded.units IS NOT NULL OR excluded.revenue IS NOT NULL OR excluded.growth IS NOT NULL THEN excluded.source
              WHEN products.units IS NOT NULL OR products.revenue IS NOT NULL OR products.growth IS NOT NULL THEN products.source
              ELSE excluded.source END,
            source_date=excluded.source_date,
            asin=CASE WHEN excluded.asin<>'' THEN excluded.asin ELSE products.asin END,
            title=excluded.title, brand=CASE WHEN excluded.brand<>'' THEN excluded.brand ELSE products.brand END,
            price=COALESCE(excluded.price,products.price), currency=COALESCE(excluded.currency,products.currency),
            units=COALESCE(excluded.units,products.units), revenue=COALESCE(excluded.revenue,products.revenue), growth=COALESCE(excluded.growth,products.growth),
            rating=COALESCE(excluded.rating,products.rating), reviews=COALESCE(excluded.reviews,products.reviews),
            bsr_rank=COALESCE(excluded.bsr_rank,products.bsr_rank),
            image_url=CASE WHEN excluded.image_url<>'' THEN excluded.image_url ELSE products.image_url END,
            product_url=CASE WHEN excluded.product_url<>'' THEN excluded.product_url ELSE products.product_url END,
            raw_json=excluded.raw_json""", row)
    con.commit(); con.close()
    return len(items)


def review_themes(title, body, category_path=""):
    text = f"{title or ''} {body or ''}".lower()
    names = LITTER_REVIEW_THEME_NAMES if "litter" in category_path.lower() or "猫砂" in category_path else REVIEW_THEMES
    return [name for name, terms in REVIEW_THEMES.items() if name in names and any(term.lower() in text for term in terms)]


def save_reviews(items):
    if not items:
        return 0
    con = connect()
    saved = 0
    for item in items:
        row = {**item, "themes_json": json.dumps(review_themes(item.get("review_title"), item.get("review_body")), ensure_ascii=False)}
        before = con.total_changes
        con.execute("""INSERT INTO reviews(country,asin,product_title,review_id,rating,review_title,review_body,review_date_text,verified,helpful_count,reviewer,review_url,source,collected_at,themes_json)
          VALUES(:country,:asin,:product_title,:review_id,:rating,:review_title,:review_body,:review_date_text,:verified,:helpful_count,:reviewer,:review_url,:source,:collected_at,:themes_json)
          ON CONFLICT(country,asin,review_id) DO UPDATE SET
            product_title=excluded.product_title,rating=excluded.rating,review_title=excluded.review_title,
            review_body=excluded.review_body,review_date_text=excluded.review_date_text,
            verified=excluded.verified,helpful_count=excluded.helpful_count,reviewer=excluded.reviewer,
            review_url=excluded.review_url,source=excluded.source,collected_at=excluded.collected_at,
            themes_json=excluded.themes_json""", row)
        saved += 1 if con.total_changes > before else 0
    con.commit(); con.close()
    return saved


def choose_review_products(country, category_path="", limit=3):
    con = connect()
    where = ["country=?", "asin<>''"]
    params = [country]
    if category_path.strip():
        where.append("category_path=?")
        params.append(category_path.strip())
    rows = [dict(row) for row in con.execute(
        "SELECT * FROM products WHERE " + " AND ".join(where) +
        " ORDER BY CASE WHEN bsr_rank IS NULL THEN 1 ELSE 0 END, bsr_rank, COALESCE(reviews,0) DESC, source_date DESC LIMIT 200",
        params,
    )]
    con.close()
    litter_category = "litter" in category_path.lower() or "猫砂" in category_path
    if litter_category:
        rows = [row for row in rows if classify_litter_product_title(row.get("title"))["status"] == "qualified"]
        positive = LITTER_AUTOMATION_TERMS
        negative = LITTER_ACCESSORY_TERMS + LITTER_MANUAL_TERMS
    else:
        positive = ("automatic", "auto ", "timed", "programmable", "wifi", "wi-fi", "camera", "cat feeder", "dog feeder", "futterautomat", "distributeur", "distributore", "dispensador")
        negative = ("water fountain", "water dispenser", "gravity", "bowl", "toy", "puzzle", "mat", "stand", "replacement", "filter")
    def score(row):
        title = str(row.get("title") or "").lower()
        return (1000 if row.get("bsr_rank") is not None else 0) + sum(30 for term in positive if term in title) - sum(40 for term in negative if term in title) + min(float(row.get("reviews") or 0), 10000) / 1000
    rows.sort(key=score, reverse=True)
    return deduplicate_products([row for row in rows if score(row) > 0], limit)


def review_summary(country="", category_path=""):
    con = connect()
    where = []
    params = []
    if country:
        where.append("r.country=?")
        params.append(country)
    if category_path:
        where.append("EXISTS(SELECT 1 FROM products p WHERE p.country=r.country AND p.asin=r.asin AND p.category_path=?)")
        params.append(category_path)
    clause = (" WHERE " + " AND ".join(where)) if where else ""
    rows = [dict(row) for row in con.execute("SELECT r.* FROM reviews r" + clause + " ORDER BY r.collected_at DESC,r.helpful_count DESC", params)]
    con.close()
    groups = {}
    for row in rows:
        group = groups.setdefault(row["country"], {"country":row["country"], "reviews":[], "themes":{}, "negative":0, "positive":0, "ratings":[]})
        group["reviews"].append(row)
        if row.get("rating") is not None:
            group["ratings"].append(row["rating"])
            if row["rating"] <= 3:
                group["negative"] += 1
            if row["rating"] >= 4:
                group["positive"] += 1
        if "litter" in category_path.lower() or "猫砂" in category_path:
            themes = review_themes(row.get("review_title"), row.get("review_body"), category_path)
        else:
            try:
                themes = json.loads(row.get("themes_json") or "[]")
            except (TypeError, ValueError):
                themes = []
        for theme in themes:
            bucket = group["themes"].setdefault(theme, {"count":0,"negative":0})
            bucket["count"] += 1
            if row.get("rating") is not None and row["rating"] <= 3:
                bucket["negative"] += 1
    result = []
    for code in COUNTRIES:
        group = groups.get(code, {"country":code,"reviews":[],"themes":{},"negative":0,"positive":0,"ratings":[]})
        count = len(group["reviews"])
        themes = sorted(({"theme":name, **values} for name, values in group["themes"].items()), key=lambda x:(x["negative"],x["count"]), reverse=True)
        result.append({
            "country": code, "label": COUNTRIES[code][0], "sample_count": count,
            "product_count": len({row["asin"] for row in group["reviews"]}),
            "average_rating": round(sum(group["ratings"])/len(group["ratings"]), 2) if group["ratings"] else None,
            "negative_count": group["negative"], "positive_count": group["positive"],
            "themes": themes, "latest_collection": max((row["collected_at"] for row in group["reviews"]), default=""),
            "reviews": group["reviews"][:100],
        })
    return result


def parse_csv_payload(content):
    text = content.decode("utf-8-sig", "replace")
    dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t") if text.strip() else csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    return list(reader)


def import_rows(filename, content):
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        rows = parse_csv_payload(content)
    elif suffix in (".xlsx", ".xlsm"):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = [str(x or "") for x in next(values, [])]
        rows = [dict(zip(headers, row)) for row in values]
    else:
        raise ValueError("支持 CSV、XLSX、XLSM 格式。旧版 XLS 请另存为 XLSX。")
    if not rows:
        raise ValueError("文件没有数据行。")
    normalized_aliases = {key: {normalize_header(x) for x in aliases} for key, aliases in HEADER_ALIASES.items()}
    mapped = []
    for raw in rows:
        cols = {normalize_header(k): v for k, v in raw.items()}
        row = {}
        for field, aliases in normalized_aliases.items():
            found = next((cols[a] for a in aliases if a in cols), None)
            row[field] = found
        country = canonical_country(row.get("country"))
        if not country or not row.get("title"):
            continue
        row["country"] = country
        row["source"] = f"导入:{Path(filename).name}"
        row["source_date"] = datetime.now(timezone.utc).isoformat()
        row["category_path"] = str(row.get("category_path") or "").strip()
        row["asin"] = str(row.get("asin") or "").strip()
        row["title"] = str(row["title"]).strip()
        row["brand"] = str(row.get("brand") or "").strip()
        row["currency"] = str(row.get("currency") or COUNTRIES[country][2]).strip()
        for number in ("price", "units", "revenue", "growth", "rating", "reviews", "bsr_rank"):
            row[number] = clean_number(row.get(number))
        row["image_url"] = str(row.get("image_url") or "").strip()
        row["product_url"] = str(row.get("product_url") or (f"https://{COUNTRIES[country][1]}/dp/{row['asin']}" if row["asin"] else "")).strip()
        mapped.append(row)
    return save_products(mapped)


def extract_features(text):
    lower = str(text or "").lower()
    found = {}
    for name, terms in FEATURES.items():
        hits = [term for term in terms if term.lower() in lower]
        found[name] = {"present": bool(hits), "evidence": hits[:3]}
    return found


def compare_product(name, specs, category_path, countries):
    our_features = extract_features(specs)
    con = connect()
    params = list(countries)
    where = "country IN (" + ",".join("?" for _ in params) + ")" if params else "1=1"
    if category_path:
        where += " AND (category_path=? OR category_path='' OR category_path IS NULL)"
        params.append(category_path)
    data = [dict(row) for row in con.execute(f"SELECT * FROM products WHERE {where} ORDER BY COALESCE(revenue,0) DESC, COALESCE(units,0) DESC, COALESCE(reviews,0) DESC LIMIT 300", params)]
    con.close()
    competitor_rows = []
    for item in data:
        features = extract_features(item["title"])
        deltas = []
        for feature, evidence in our_features.items():
            if evidence["present"] and not features[feature]["present"]:
                verdict = "我方描述有、竞品标题未见证据，需打开详情页核实"
            elif evidence["present"] and features[feature]["present"]:
                verdict = "双方标题/资料均有，非独占优势"
            elif not evidence["present"] and features[feature]["present"]:
                verdict = "竞品标题有、我方资料未提及，待确认是否缺失"
            else:
                verdict = "双方资料均未提及"
            deltas.append({"feature": feature, "verdict": verdict, "our_evidence": evidence["evidence"], "competitor_evidence": features[feature]["evidence"]})
        competitor_rows.append({**item, "feature_comparison": deltas})
    supported_advantages = [name for name, evidence in our_features.items() if evidence["present"] and any(not extract_features(row["title"])[name]["present"] for row in data)]
    return {"product_name": name, "category_path": category_path, "features": our_features, "competitors": competitor_rows, "supported_differentiators": supported_advantages, "limitations": ["竞品标题没写不代表商品没有此功能，差异结论需逐条打开链接核实。", "销量、销售额和增长率仅在数据源提供时参与机会判断；不会用公开页面数据或模型补造。", "页面截图/图片对比受图片角度、光线和尺寸影响，需人工确认产品结构。"], "created_at": datetime.now(timezone.utc).isoformat()}


def opportunity_tags(rows):
    scoped = [r for r in rows if r.get("country") and r.get("category_path")]
    by_scope = {}
    for row in scoped:
        by_scope.setdefault((row["country"], row["category_path"]), []).append(row)
    tagged = []
    for row in rows:
        group = by_scope.get((row.get("country"), row.get("category_path")), [])
        units = sorted([x["units"] for x in group if x["units"] is not None])
        revenues = sorted([x["revenue"] for x in group if x["revenue"] is not None])
        prices = sorted([x["price"] for x in group if x["price"] is not None])
        q_units = units[int((len(units)-1)*0.75)] if units else None
        q_revenue = revenues[int((len(revenues)-1)*0.75)] if revenues else None
        q_price = prices[int((len(prices)-1)*0.75)] if prices else None
        tags = []
        if row.get("growth") is not None and row["growth"] >= 20:
            tags.append("高增长")
        if (row.get("units") is not None and q_units is not None and row["units"] >= q_units) or (row.get("revenue") is not None and q_revenue is not None and row["revenue"] >= q_revenue):
            tags.append("流量款")
        if row.get("price") is not None and q_price is not None and row["price"] >= q_price:
            tags.append("利润款候选")
        tagged.append({**row, "opportunity_tags": tags, "thresholds": {"units_p75": q_units, "revenue_p75": q_revenue, "price_p75": q_price}, "opportunity_note": "样本内相对筛选，不等同于净利润；价格需结合成本、广告、佣金、退货和税费。"})
    return tagged


def export_xlsx(category_path=""):
    con = connect()
    sql = "SELECT * FROM products" + (" WHERE category_path=? OR category_path=''" if category_path else "") + " ORDER BY country,category_path,CASE WHEN bsr_rank IS NULL THEN 1 ELSE 0 END,bsr_rank,COALESCE(revenue,0) DESC,COALESCE(units,0) DESC"
    data = [dict(x) for x in con.execute(sql, (category_path,) if category_path else ())]
    analyses = [dict(x) for x in con.execute("SELECT product_name,category_path,created_at,result_json FROM analyses ORDER BY created_at DESC LIMIT 50")]
    reviews = [dict(x) for x in con.execute("SELECT * FROM reviews ORDER BY country,collected_at DESC,helpful_count DESC")]
    histories = [dict(x) for x in con.execute("SELECT * FROM analysis_runs ORDER BY started_at DESC LIMIT 100")]
    con.close()
    data = opportunity_tags(data)
    book = Workbook(); products = book.active; products.title = "七国商品数据"
    headers = ["国家", "三级类目", "数据源", "采集时间", "ASIN", "BSR排名", "品牌", "标题", "价格", "币种", "月销量", "月销售额", "增长率(%)", "评分", "评论数", "机会标签", "商品图", "商品链接"]
    products.append(headers)
    for row in data:
        products.append([COUNTRIES.get(row["country"], (row["country"],))[0], row["category_path"], row["source"], row["source_date"], row["asin"], row["bsr_rank"], row["brand"], row["title"], row["price"], row["currency"], row["units"], row["revenue"], row["growth"], row["rating"], row["reviews"], "、".join(row["opportunity_tags"]), row["image_url"], row["product_url"]])
    for ws in book.worksheets:
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="17324D")
        for cells in ws.columns:
            ws.column_dimensions[cells[0].column_letter].width = min(48, max(12, max(len(str(c.value or "")) for c in cells[:100]) + 2))
    opp = book.create_sheet("机会定义")
    opp.append(["标签", "可复核定义", "所需字段"])
    opp.append(["高增长", "数据源增长率≥20%", "growth；需注明周期和来源"])
    opp.append(["流量款", "同国家+三级类目样本的月销量或月销售额达到P75", "units 或 revenue"])
    opp.append(["利润款候选", "同国家+三级类目样本售价达到P75；不代表净利高", "price；还需成本/费用判断利润"])
    feat = book.create_sheet("功能对照")
    feat.append(["产品名", "三级类目", "分析时间", "功能", "判定", "我方证据", "竞品证据", "ASIN"])
    for analysis in analyses:
        result = json.loads(analysis["result_json"])
        for item in result.get("competitors", []):
            for diff in item.get("feature_comparison", []):
                feat.append([analysis["product_name"], analysis["category_path"], analysis["created_at"], diff["feature"], diff["verdict"], ", ".join(diff["our_evidence"]), ", ".join(diff["competitor_evidence"]), item.get("asin")])
    review_sheet = book.create_sheet("七国评论原文")
    review_sheet.append(["国家","ASIN","商品标题","评论ID","星级","评论标题","评论正文","评论日期原文","已验证购买","有用数","评论者","评论链接","来源","采集时间","主题"])
    for row in reviews:
        review_sheet.append([COUNTRIES.get(row["country"],(row["country"],))[0],row["asin"],row["product_title"],row["review_id"],row["rating"],row["review_title"],row["review_body"],row["review_date_text"],"是" if row["verified"] else "否",row["helpful_count"],row["reviewer"],row["review_url"],row["source"],row["collected_at"],"、".join(json.loads(row["themes_json"] or "[]"))])
    insight_sheet = book.create_sheet("七国评论分析")
    insight_sheet.append(["国家","评论样本","商品数","平均星级","低星评论","高星评论","主要主题","最近采集时间"])
    for item in review_summary(category_path=category_path):
        insight_sheet.append([item["label"],item["sample_count"],item["product_count"],item["average_rating"],item["negative_count"],item["positive_count"],"；".join(f'{theme["theme"]}:{theme["count"]}条/低星{theme["negative"]}条' for theme in item["themes"][:8]),item["latest_collection"]])
    history_sheet = book.create_sheet("分析历史")
    history_sheet.append(["任务ID","类型","产品名","三级类目","国家","来源月份","样本数","开始时间","完成时间","状态","结论","消息"])
    for row in histories:
        history_sheet.append([row["id"],row["kind"],row["product_name"],row["category_path"],row["countries"],row["source_period"],row["sample_count"],row["started_at"],row["finished_at"],row["status"],row["conclusions_json"],row["message"]])
    for ws in (review_sheet, insight_sheet, history_sheet):
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="17324D")
        for cells in ws.columns:
            ws.column_dimensions[cells[0].column_letter].width = min(56, max(12, max(len(str(c.value or "")) for c in cells[:100]) + 2))
    output = io.BytesIO(); book.save(output); return output.getvalue()


def json_bytes(value):
    return json.dumps(value, ensure_ascii=False, default=str).encode("utf-8")


def connector_snapshot(run_id):
    with _CONNECTOR_LOCK:
        return dict(_CONNECTOR_RUNS.get(run_id, {}))


def update_connector(run_id, **values):
    with _CONNECTOR_LOCK:
        run = _CONNECTOR_RUNS.setdefault(run_id, {})
        run.update(values)
        snapshot = dict(run)
    con = connect()
    con.execute(
        "UPDATE connector_runs SET status=?,finished_at=?,message=?,result_json=? WHERE id=?",
        (snapshot.get("status", "RUNNING"), snapshot.get("finished_at"), snapshot.get("message", ""), json.dumps(snapshot.get("result", {}), ensure_ascii=False), run_id),
    )
    con.commit(); con.close()


def active_connector_run():
    with _CONNECTOR_LOCK:
        return next((dict(run) for run in _CONNECTOR_RUNS.values() if run.get("status") == "RUNNING"), None)


def create_connector_run(kind, target):
    if active_connector_run():
        raise RuntimeError("已有数据连接任务正在运行，请等待完成或关闭其浏览器窗口。")
    run_id = uuid.uuid4().hex
    started_at = datetime.now(timezone.utc).isoformat()
    run = {"id":run_id,"kind":kind,"status":"RUNNING","started_at":started_at,"message":"准备启动","page":0,"count":0,"result":{}}
    with _CONNECTOR_LOCK:
        _CONNECTOR_RUNS[run_id] = run
    con=connect(); con.execute("INSERT INTO connector_runs(id,kind,status,started_at,message,result_json) VALUES(?,?,?,?,?,?)",(run_id,kind,"RUNNING",started_at,"准备启动","{}")); con.commit(); con.close()
    threading.Thread(target=target,args=(run_id,),daemon=True).start()
    return run


def execute_sellersprite_login(run_id):
    def progress(message, page, count):
        update_connector(run_id,message=message,page=page,count=count)
    try:
        result=sellersprite_open_login(SELLERSPRITE_PROFILE_DIR,progress)
        update_connector(run_id,status="COMPLETED",finished_at=datetime.now(timezone.utc).isoformat(),message="卖家精灵登录会话已保存在本机",result=result)
    except Exception as exc:
        update_connector(run_id,status="FAILED",finished_at=datetime.now(timezone.utc).isoformat(),message=str(exc),result={"error":str(exc)})


def execute_sellersprite_sync(run_id, url, category_path, max_products):
    def progress(message, page, count):
        update_connector(run_id,message=message,page=page,count=count)
    try:
        rows=sellersprite_sync_results(url,category_path,SELLERSPRITE_PROFILE_DIR,max_products,progress)
        saved=save_products(rows)
        countries=sorted({row["country"] for row in rows})
        result={"saved":saved,"countries":countries,"unique_products":len(rows)}
        update_connector(run_id,status="COMPLETED",finished_at=datetime.now(timezone.utc).isoformat(),message=f"同步完成，保存 {saved} 个去重商品",count=len(rows),result=result)
    except Exception as exc:
        update_connector(run_id,status="FAILED",finished_at=datetime.now(timezone.utc).isoformat(),message=str(exc),result={"error":str(exc)})


def execute_review_sync(run_id, countries, category_path, products_per_country, reviews_per_product):
    started_at = datetime.now(timezone.utc).isoformat()
    analysis_id = uuid.uuid4().hex
    con = connect()
    con.execute("INSERT INTO analysis_runs(id,kind,category_path,countries,started_at,status,message) VALUES(?,?,?,?,?,?,?)", (analysis_id,"amazon_reviews",category_path,json.dumps(countries),started_at,"RUNNING","准备采集七国评论"))
    con.commit(); con.close()
    outcomes = []
    total = 0
    try:
        for country in countries:
            chosen = choose_review_products(country, category_path, products_per_country)
            outcome = {"country":country,"label":COUNTRIES[country][0],"products":len(chosen),"reviews":0,"status":"RUNNING","selection":"BSR优先；无BSR时按评论量"}
            outcomes.append(outcome)
            if not chosen:
                outcome.update({"status":"failed","error":"同类目没有可用 ASIN"})
                update_connector(run_id,message=f"{outcome['label']}：没有可用 ASIN",result={"countries":outcomes,"saved":total})
                continue
            for index, product in enumerate(chosen, start=1):
                def progress(message, page, count):
                    update_connector(run_id,message=message,page=page,count=total+count,result={"countries":outcomes,"saved":total})
                try:
                    rows = amazon_collect_reviews(country, COUNTRIES[country][1], product["asin"], product["title"], AMAZON_PROFILE_DIR, RAW_DIR / "reviews", reviews_per_product, progress)
                    saved = save_reviews(rows)
                    total += saved
                    outcome["reviews"] += saved
                    outcome["last_asin"] = product["asin"]
                    update_connector(run_id,message=f"{outcome['label']}：完成 {index}/{len(chosen)} 个商品，累计 {outcome['reviews']} 条评论",count=total,result={"countries":outcomes,"saved":total})
                except Exception as exc:
                    outcome["warning"] = str(exc)[:500]
            outcome["status"] = "ok" if outcome["reviews"] else "failed"
        summaries = review_summary(category_path=category_path)
        conclusions = {item["country"]:{"sample_count":item["sample_count"],"product_count":item["product_count"],"average_rating":item["average_rating"],"negative_count":item["negative_count"],"top_themes":item["themes"][:5]} for item in summaries if item["country"] in countries}
        finished = datetime.now(timezone.utc).isoformat()
        status = "COMPLETED" if total else "FAILED"
        message = f"评论采集完成，保存 {total} 条真实评论" if total else "未取得评论；请查看各国家错误信息"
        con = connect(); con.execute("UPDATE analysis_runs SET sample_count=?,finished_at=?,status=?,conclusions_json=?,message=? WHERE id=?",(total,finished,status,json.dumps(conclusions,ensure_ascii=False),message,analysis_id)); con.commit(); con.close()
        update_connector(run_id,status=status,finished_at=finished,message=message,count=total,result={"countries":outcomes,"saved":total,"analysis_id":analysis_id,"conclusions":conclusions})
    except Exception as exc:
        finished = datetime.now(timezone.utc).isoformat()
        con = connect(); con.execute("UPDATE analysis_runs SET finished_at=?,status='FAILED',message=? WHERE id=?",(finished,str(exc),analysis_id)); con.commit(); con.close()
        update_connector(run_id,status="FAILED",finished_at=finished,message=str(exc),result={"error":str(exc),"countries":outcomes,"saved":total})


class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        print(f"[{self.log_date_time_string()}] {fmt % args}")

    def send_json(self, value, status=200):
        body = json_bytes(value); self.send_response(status); self.send_header("Content-Type", "application/json; charset=utf-8"); self.send_header("Content-Length", str(len(body))); self.send_header("Cache-Control", "no-store"); self.end_headers(); self.wfile.write(body)

    def body(self):
        return self.rfile.read(int(self.headers.get("Content-Length", "0")))

    def do_GET(self):
        path = urllib.parse.urlparse(self.path).path
        if path == "/api/status":
            con=connect(); count=con.execute("SELECT COUNT(*) FROM products").fetchone()[0]; countries=con.execute("SELECT COUNT(DISTINCT country) FROM products").fetchone()[0]; latest=con.execute("SELECT MAX(source_date) FROM products").fetchone()[0]; review_count=con.execute("SELECT COUNT(*) FROM reviews").fetchone()[0]; analysis_count=con.execute("SELECT COUNT(*) FROM analysis_runs").fetchone()[0]; con.close()
            return self.send_json({"app":"AM master","data_dir":str(DATA_DIR),"products":count,"countries_with_data":countries,"latest_update":latest,"reviews":review_count,"analysis_runs":analysis_count,"tokenless":True})
        if path == "/api/products":
            q=urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query); country=q.get("country",[""])[0]; category=q.get("category",[""])[0]
            con=connect(); where=[]; args=[]
            if country: where.append("country=?"); args.append(country)
            if category: where.append("category_path=?"); args.append(category)
            sql="SELECT * FROM products"+(" WHERE "+" AND ".join(where) if where else "")+" ORDER BY CASE WHEN bsr_rank IS NULL THEN 1 ELSE 0 END,bsr_rank,source_date DESC LIMIT 1000"
            rows=[dict(x) for x in con.execute(sql,args)]; con.close(); return self.send_json({"items":opportunity_tags(rows)})
        if path == "/api/categories":
            con=connect(); values=[r[0] for r in con.execute("SELECT DISTINCT category_path FROM products WHERE category_path<>'' ORDER BY category_path")]; con.close(); return self.send_json({"items":values})
        if path == "/api/sellersprite/status":
            return self.send_json({"profile_ready":(SELLERSPRITE_PROFILE_DIR / ".logged_in").exists(),"profile_dir":str(SELLERSPRITE_PROFILE_DIR),"active_run":active_connector_run()})
        if path == "/api/sellersprite-run":
            run_id=urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query).get("id",[""])[0]
            run=connector_snapshot(run_id)
            if not run: return self.send_json({"error":"找不到该卖家精灵任务"},404)
            return self.send_json(run)
        if path == "/api/reviews":
            q=urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query); country=q.get("country",[""])[0]; category=q.get("category",[""])[0]
            return self.send_json({"items":review_summary(country,category)})
        if path == "/api/analysis-history":
            con=connect(); rows=[dict(x) for x in con.execute("SELECT * FROM analysis_runs ORDER BY started_at DESC LIMIT 100")]; con.close()
            for row in rows:
                for field in ("countries","conclusions_json"):
                    try: row[field]=json.loads(row[field] or ("[]" if field=="countries" else "{}"))
                    except (TypeError,ValueError): pass
            return self.send_json({"items":rows})
        if path == "/api/search-run":
            run_id=urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query).get("id",[""])[0]
            con=connect(); row=con.execute("SELECT * FROM searches WHERE id=?",(run_id,)).fetchone(); con.close()
            if not row: return self.send_json({"error":"找不到该搜索任务"},404)
            result=dict(row)
            try: result["progress"]=json.loads(result["message"] or "[]")
            except json.JSONDecodeError: result["progress"]=[]
            return self.send_json(result)
        if path == "/" or path.startswith("/static/"):
            relative="static/index.html" if path == "/" else path.lstrip("/")
            file=(ROOT / relative).resolve()
            if STATIC not in file.parents: return self.send_error(404)
            if not file.is_file(): return self.send_error(404)
            body=file.read_bytes(); kind="text/html; charset=utf-8" if file.suffix==".html" else "text/css; charset=utf-8" if file.suffix==".css" else "text/csv; charset=utf-8" if file.suffix==".csv" else "application/javascript; charset=utf-8"
            self.send_response(200); self.send_header("Content-Type",kind); self.send_header("Content-Length",str(len(body))); self.send_header("Cache-Control","no-store"); self.end_headers(); self.wfile.write(body); return
        return self.send_error(404)

    def do_POST(self):
        path=urllib.parse.urlparse(self.path).path
        try:
            if path == "/api/import":
                content=self.body(); message=BytesParser(policy=default).parsebytes((f"Content-Type: {self.headers.get('Content-Type','application/octet-stream')}\r\nMIME-Version: 1.0\r\n\r\n".encode()+content))
                filename=urllib.parse.unquote(self.headers.get("X-File-Name", "import.csv")); payload=content
                if message.is_multipart():
                    part=next((p for p in message.iter_parts() if p.get_filename()),None)
                    if not part: return self.send_json({"error":"未找到上传文件"},400)
                    filename=part.get_filename(); payload=part.get_payload(decode=True)
                count=import_rows(filename,payload); return self.send_json({"imported":count})
            data=json.loads(self.body() or b"{}")
            if path == "/api/search":
                keyword=str(data.get("keyword","")).strip(); category=str(data.get("category_path","")).strip(); countries=[x for x in data.get("countries",list(COUNTRIES)) if x in COUNTRIES]
                if not keyword: return self.send_json({"error":"请输入产品名或关键词"},400)
                if not countries: return self.send_json({"error":"至少选择一个目标国家"},400)
                run_id=uuid.uuid4().hex; started=datetime.now(timezone.utc).isoformat(); con=connect(); con.execute("INSERT INTO searches(id,keyword,category_path,countries,status,started_at) VALUES(?,?,?,?,?,?)",(run_id,keyword,category,json.dumps(countries),"RUNNING",started)); con.commit(); con.close()
                thread=threading.Thread(target=execute_search,args=(run_id,keyword,category,countries),daemon=True); thread.start()
                return self.send_json({"run_id":run_id,"status":"RUNNING","countries":countries})
            if path == "/api/analyze":
                name=str(data.get("product_name","")).strip(); specs=str(data.get("product_text","")).strip(); category=str(data.get("category_path","")).strip(); countries=[x for x in data.get("countries",list(COUNTRIES)) if x in COUNTRIES]
                if not name or not specs: return self.send_json({"error":"产品名称和产品资料都必填"},400)
                if not countries: return self.send_json({"error":"至少选择一个对照市场"},400)
                result=compare_product(name,specs,category,countries); analysis_id=uuid.uuid4().hex; con=connect(); con.execute("INSERT INTO analyses VALUES(?,?,?,?,?,?)",(analysis_id,name,category,specs,result["created_at"],json.dumps(result,ensure_ascii=False))); con.execute("INSERT INTO analysis_runs(id,kind,product_name,category_path,countries,sample_count,started_at,finished_at,status,conclusions_json,message) VALUES(?,?,?,?,?,?,?,?,?,?,?)",(uuid.uuid4().hex,"product_comparison",name,category,json.dumps(countries),len(result.get("competitors",[])),result["created_at"],result["created_at"],"COMPLETED",json.dumps({"supported_differentiators":result.get("supported_differentiators",[]),"limitations":result.get("limitations",[])},ensure_ascii=False),"本地规则功能对照完成")); con.commit(); con.close(); return self.send_json({"id":analysis_id,**result})
            if path == "/api/sellersprite/login":
                run=create_connector_run("sellersprite_login",execute_sellersprite_login)
                return self.send_json(run,202)
            if path == "/api/sellersprite/sync":
                url=str(data.get("url","")).strip(); category=str(data.get("category_path","")).strip(); maximum=int(data.get("max_products",100) or 100)
                validate_sellersprite_url(url)
                run=create_connector_run("sellersprite_sync",lambda run_id: execute_sellersprite_sync(run_id,url,category,maximum))
                return self.send_json(run,202)
            if path == "/api/reviews/sync":
                countries=[x for x in data.get("countries",list(COUNTRIES)) if x in COUNTRIES]; category=str(data.get("category_path","")).strip(); product_limit=max(1,min(int(data.get("products_per_country",3) or 3),10)); review_limit=max(1,min(int(data.get("reviews_per_product",10) or 10),50))
                if not countries: return self.send_json({"error":"至少选择一个目标国家"},400)
                run=create_connector_run("amazon_reviews",lambda run_id: execute_review_sync(run_id,countries,category,product_limit,review_limit))
                return self.send_json(run,202)
            return self.send_error(404)
        except Exception as exc:
            return self.send_json({"error":str(exc)},500)

class ExportHandler(Handler):
    def do_GET(self):
        if urllib.parse.urlparse(self.path).path == "/api/export.xlsx":
            try:
                q=urllib.parse.parse_qs(urllib.parse.urlparse(self.path).query); payload=export_xlsx(q.get("category",[""])[0]); self.send_response(200); self.send_header("Content-Type","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"); self.send_header("Content-Disposition","attachment; filename=AM-master-market-analysis.xlsx"); self.send_header("Content-Length",str(len(payload))); self.end_headers(); self.wfile.write(payload)
            except Exception as exc: self.send_json({"error":str(exc)},500)
            return
        super().do_GET()


def execute_search(run_id, keyword, category, countries):
    outcomes=[]; saved=0
    for code in countries:
        outcome={"country":code,"label":COUNTRIES[code][0],"status":"RUNNING","page":0,"count":0,"target":MAX_PRODUCTS_PER_MARKET}
        outcomes.append(outcome)
        def publish_progress(page, count, added, warning=None):
            outcome.update({"page":page,"count":count,"added":added})
            if warning:
                outcome["warning"] = warning
            con=connect(); con.execute("UPDATE searches SET message=? WHERE id=?",(json.dumps({"countries":outcomes,"saved":saved},ensure_ascii=False),run_id)); con.commit(); con.close()
        bsr_items=load_bsr_category_products(code,category,MAX_PRODUCTS_PER_MARKET)
        if bsr_items:
            publish_progress(0,len(bsr_items),len(bsr_items),f"已优先载入 {len(bsr_items)} 条 BSR 数据，正在用搜索页补足 {MAX_PRODUCTS_PER_MARKET-len(bsr_items)} 条")
        elif not category:
            publish_progress(0,0,0,"三级类目为空，无法可靠匹配 BSR；本次使用搜索页数据")
        else:
            publish_progress(0,0,0,"当前国家和三级类目暂无 BSR 数据，正在使用搜索页数据")
        amazon_error = ""
        search_items=[]
        if len(bsr_items)<MAX_PRODUCTS_PER_MARKET:
            def publish_search_progress(page, count, added, warning=None):
                merged_count=len(merge_bsr_then_search(bsr_items,[{}] * 0,MAX_PRODUCTS_PER_MARKET)) + count
                publish_progress(page,min(MAX_PRODUCTS_PER_MARKET,merged_count),added,warning)
            try:
                search_items=amazon_search(code,keyword,category,publish_search_progress)
            except Exception as exc:
                amazon_error=str(exc)[:500]
        items=merge_bsr_then_search(bsr_items,search_items,MAX_PRODUCTS_PER_MARKET)
        if items:
            saved+=save_products(items); outcome.update({"status":"ok" if len(items)>=MAX_PRODUCTS_PER_MARKET else "partial","count":len(items)})
            if len(items)<MAX_PRODUCTS_PER_MARKET:
                detail=amazon_error or outcome.get("warning") or "可用来源商品不足"
                outcome["warning"] = f"{detail}；最终取得 {len(items)}/{MAX_PRODUCTS_PER_MARKET} 条"
        else:
            outcome.update({"status":"failed","error":amazon_error or "Amazon 与同类目授权缓存均无可用商品"})
        con=connect(); con.execute("UPDATE searches SET message=? WHERE id=?",(json.dumps({"countries":outcomes,"saved":saved},ensure_ascii=False),run_id)); con.commit(); con.close()
        time.sleep(0.6)
    status="COMPLETED" if any(x["status"] in ("ok","partial") for x in outcomes) else "FAILED"
    con=connect(); con.execute("UPDATE searches SET status=?,finished_at=?,message=? WHERE id=?",(status,datetime.now(timezone.utc).isoformat(),json.dumps({"countries":outcomes,"saved":saved},ensure_ascii=False),run_id)); con.commit(); con.close()


if __name__ == "__main__":
    connect().close()
    server=ThreadingHTTPServer(("127.0.0.1",PORT),ExportHandler)
    print(f"AM master running at http://127.0.0.1:{PORT} | data: {DATA_DIR}",flush=True)
    server.serve_forever()
